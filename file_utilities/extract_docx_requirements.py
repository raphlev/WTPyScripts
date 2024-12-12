import sys
import os
from docx import Document
from openpyxl import Workbook
import re

############################################
#              PARAMETRES GLOBAUX
############################################

# Liste des styles de titre (headings) considérés comme chapitres
CHAPTER_STYLES = ["Heading 1", "Heading 2", "Heading 3", "Heading 4", "Heading 5"]

# Expression régulière pour détecter la ligne annonçant les exigences
EXIGENCES_TRIGGER_REGEX = r"^Exigences?\s*:"

# Noms des colonnes dans le fichier Excel
EXCEL_COLUMNS = ["Titre", "Nom Exigence", "Description"]

# Regex pour détecter les paragraphes à puces.
# Ajustez les caractères selon les puces réellement utilisées dans votre document.
BULLET_REGEX = r"^[•-]\s*"

############################################
#             FONCTION PRINCIPALE
############################################

def extract_requirements(docx_path, xlsx_path):
    # Ouvrir le document Word
    document = Document(docx_path)
    
    # Préparer une structure pour stocker (Chapitre, Liste d'exigences)
    # format: { "Chapitre": ["exigence1", "exigence2", ...], ... }
    chapter_requirements = {}
    
    current_chapter = None
    collecting = False  # Indique si on est en train de collecter des exigences
    current_requirement_text = None  # Pour stocker l'exigence en cours de construction

    # Parcourir tous les paragraphes du document
    for para in document.paragraphs:
        style_name = para.style.name if para.style else ""
        text = para.text.strip()

        # Vérifier si le paragraphe est un chapitre
        if style_name in CHAPTER_STYLES:
            # Nouveau chapitre détecté : on arrête la collecte d'exigences en cours
            # et on enregistre l'exigence en cours s'il y en a une
            if current_requirement_text and current_chapter is not None:
                chapter_requirements[current_chapter].append(current_requirement_text)
                current_requirement_text = None
            
            collecting = False
            current_chapter = text
            if current_chapter not in chapter_requirements:
                chapter_requirements[current_chapter] = []
        
        else:
            # Paragraphe "normal" (corps)
            if not text:
                continue
            
            # Si on détecte la ligne déclenchant la collecte
            if re.match(EXIGENCES_TRIGGER_REGEX, text, re.IGNORECASE):
                # Avant de commencer à collecter, on clôture l'exigence en cours
                if current_requirement_text and current_chapter is not None:
                    chapter_requirements[current_chapter].append(current_requirement_text)
                    current_requirement_text = None
                collecting = True
                continue
            
            # Si on est en mode collecte
            if collecting and current_chapter is not None:
                # On détermine si le paragraphe est une puce
                if re.match(BULLET_REGEX, text):
                    # C'est une puce : on l'ajoute à l'exigence en cours
                    if current_requirement_text is None:
                        # Si aucune exigence n'était encore ouverte, on en crée une
                        current_requirement_text = text
                    else:
                        # Sinon, on ajoute simplement ce paragraphe à la fin de l'exigence courante
                        current_requirement_text += "\n" + text
                else:
                    # Paragraphe normal, non bullet
                    # On clôture l'exigence précédente si elle existe
                    if current_requirement_text is not None:
                        chapter_requirements[current_chapter].append(current_requirement_text)
                    # On démarre une nouvelle exigence
                    current_requirement_text = text

    # A la fin du document, si une exigence est en cours, on l'ajoute
    if current_requirement_text and current_chapter is not None:
        chapter_requirements[current_chapter].append(current_requirement_text)
    
    # Génération du fichier Excel
    # On écrase le fichier si déjà existant
    if os.path.exists(xlsx_path):
        os.remove(xlsx_path)
    
    wb = Workbook()
    ws = wb.active
    
    # Ecrire les en-têtes
    ws.append(EXCEL_COLUMNS)
    
    # Définir une fonction locale pour générer le nom d'exigence
    # Ici, on prend la partie de l'exigence avant le ":" si présent
    def extract_req_name_from_text(requirement_text):
        parts = requirement_text.split(":", 1)
        return parts[0].strip() if len(parts) > 1 else requirement_text.strip()
    
    # Remplir le tableau
    for chapter, req_list in chapter_requirements.items():
        # Extraire uniquement la partie du chapitre avant le premier ":" s'il existe
        chapter_title = chapter.split(":", 1)[0].strip()
        
        for req in req_list:
            req_name = extract_req_name_from_text(req)
            ws.append([chapter_title, req_name, req])
    
    # Enregistrer le fichier Excel
    wb.save(xlsx_path)

############################################
#                  MAIN
############################################

if __name__ == "__main__":
    # Exemple d'utilisation:
    # python extract_docx_requirements.py input.docx output.xlsx
    if len(sys.argv) < 3:
        print("Usage: python script.py input.docx output.xlsx")
        sys.exit(1)
    docx_path = sys.argv[1]
    xlsx_path = sys.argv[2]
    extract_requirements(docx_path, xlsx_path)
    print(f"Fichier Excel généré : {xlsx_path}")
